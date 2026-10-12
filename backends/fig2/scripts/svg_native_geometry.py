"""Resolve SVG geometry before writing editable DrawingML paths.

Cairo resolves CSS, viewports, text metrics and nested transforms. Skia PathOps
intersects real Bezier geometry with clips and expands strokes. Gradients become
bounded vector color bands, never a first-stop solid or a raster image.
Text retains Unicode in native text boxes, placed from the resolved SVG baselines.
"""
import math
import re
from pathlib import Path
import xml.etree.ElementTree as ET

import cairosvg
import cairocffi as cairo
import pathops
from fontTools.svgLib.path import parse_path
from svgelements import Matrix, Color
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement

from display_check import DisplayError, canvas

NS = 'http://www.w3.org/2000/svg'
IDENTITY = (1., 0., 0., 1., 0., 0.)


def multiply(a, b):
    # Column vectors: parent @ local, not the reverse order.
    A,B,C,D,E,F = a; g,h,i,j,k,l = b
    return (A*g+C*h,B*g+D*h,A*i+C*j,B*i+D*j,A*k+C*l+E,B*k+D*l+F)


def matrix(value):
    m = Matrix(value) if value else Matrix()
    return (m.a,m.b,m.c,m.d,m.e,m.f)


def rectangle(x, y, w, h):
    p=pathops.Path();p.moveTo(x,y);p.lineTo(x+w,y);p.lineTo(x+w,y+h);p.lineTo(x,y+h);p.close()
    return p


def intersect(a, b):
    a=pathops.Path(a);b=pathops.Path(b)
    a.convertConicsToQuads(.005);b.convertConicsToQuads(.005)
    return pathops.op(a,b,pathops.PathOp.INTERSECTION,fix_winding=True)


def geometry(n):
    tag=n.tag.split('}')[-1]
    if tag=='path':
        p=pathops.Path();parse_path(n.get('d',''),p.getPen())
    elif tag=='rect':
        p=rectangle(*[float(n.get(k,0)) for k in ('x','y','width','height')])
    else: raise DisplayError('Unresolved vector geometry: '+tag)
    if n.get('fill-rule',n.get('clip-rule'))=='evenodd':p.fillType=pathops.FillType.EVEN_ODD
    return p


def color(value):
    c=Color(value)
    if c.value is None:raise DisplayError('Unresolved color: '+value)
    return (c.red/255,c.green/255,c.blue/255,c.alpha/255)


def resolve_font_fallback(root):
    """Cairo toy text has no OS fallback; resolve missing glyphs before outlining."""
    from illustrator_svg import styles
    computed=styles(root);cache={};changes=0
    ctx=cairo.Context(cairo.ImageSurface(cairo.FORMAT_ARGB32,1,1))
    def supports(family,ch):
        key=(family,ch)
        if key not in cache:
            ctx.select_font_face(family)
            cache[key]=all(g[0] for g in ctx.get_scaled_font().text_to_glyphs(0,0,ch,False))
        return cache[key]
    def runs(text,family):
        nonlocal changes
        found=[]
        for ch in text:
            chosen=family
            if not ch.isspace() and not supports(family,ch):
                chosen=next((f for f in ('SimSun','Microsoft YaHei','Arial','DejaVu Sans') if supports(f,ch)),None)
                if chosen is None:raise DisplayError('A text glyph has no installed font; install the requested font before conversion')
                changes+=1
            if found and found[-1][0]==chosen:found[-1]=(chosen,found[-1][1]+ch)
            else:found.append((chosen,ch))
        return found
    def visit(n,parent_family='Times New Roman'):
        family=computed.get(n,{}).get('font-family',parent_family).split(',')[0].strip(' \"\'')
        children=list(n)
        for item in children:visit(item,family)
        if n.tag.split('}')[-1] not in ('text','tspan'):return
        # Rebuild only chunks with missing glyphs. Keep positioning and original spans.
        if n.text:
            chunks=runs(n.text,family)
            if any(f!=family for f,_ in chunks):
                n.text=None
                for i,(f,t) in enumerate(chunks):
                    span=ET.Element('{'+NS+'}tspan',{'style':'font-family:'+f+' !important'});span.text=t;n.insert(i,span)
        for item in children:
            if item.tail:
                chunks=runs(item.tail,family)
                if any(f!=family for f,_ in chunks):
                    item.tail=None;index=list(n).index(item)+1
                    for f,t in chunks:
                        span=ET.Element('{'+NS+'}tspan',{'style':'font-family:'+f+' !important'});span.text=t;n.insert(index,span);index+=1
    visit(root)
    return changes


def resolve_baseline_shifts(root):
    """Cairo ignores baseline-shift; translate only the affected text paint node."""
    from illustrator_svg import styles
    computed=styles(root);count=0
    def length(v,em):
        v=str(v).strip()
        if v.endswith('%'):return float(v[:-1])*em/100
        if v.endswith('em'):return float(v[:-2])*em
        return float(v.removesuffix('px').removesuffix('pt'))*(96/72 if v.endswith('pt') else 1)
    def visit(n,parent_size=16):
        nonlocal count
        style=computed.get(n,{})
        size=length(style.get('font-size',parent_size),parent_size)
        shift=style.get('baseline-shift','baseline')
        if n.tag.split('}')[-1] in ('text','tspan') and shift not in ('baseline','0','0px'):
            # Named shifts use conventional em offsets and are flagged for visual review.
            dy=size*.2 if shift=='sub' else -size*.4 if shift=='super' else -length(shift,size)
            n.set('transform','translate(0 '+str(dy)+') '+n.get('transform',''))
            count+=1
        for child in n:visit(child,size)
    visit(root)
    return count


def native_paths(data, *, gradient_steps=128):
    source=ET.fromstring(data)
    # CairoSVG does not implement general SVG filters. Do not quietly erase them.
    if any(n.tag.split('}')[-1] in ('filter','image','foreignObject') for n in source.iter()):
        raise DisplayError('Source has a filter/image requiring explicit vector preparation; original SVG retained')
    baseline_adjustments=resolve_baseline_shifts(source)
    font_fallbacks=resolve_font_fallback(source)
    from native_text import capture, TextMatcher
    normalized,text_records=capture(ET.tostring(source))
    text_matcher=TextMatcher(text_records)
    root=ET.fromstring(normalized);dims=canvas(normalized)
    ids={n.get('id'):n for n in root.iter() if n.get('id')}
    rows=[];stats={'clipped_paths':0,'gradient_objects':0,'outlined_text':False,'group_opacity_approximations':0,'font_fallback_glyphs':font_fallbacks,'baseline_adjustments':baseline_adjustments}
    viewport=rectangle(0,0,dims['width_px'],dims['height_px'])

    def append(p, rgba, clip):
        if not len(p) or rgba[3]<=0:return
        if clip is not None:
            p=intersect(p,clip);stats['clipped_paths']+=1
        if len(p):rows.append((pathops.simplify(p,fix_winding=True),rgba))

    def gradient(p, name, transform, opacity, clip):
        g=ids[name];gm=multiply(transform,matrix(g.get('gradientTransform')))
        stops=[]
        for s in g:
            offset=s.get('offset','0');offset=float(offset[:-1])/100 if offset.endswith('%') else float(offset)
            rgba=color(s.get('stop-color','black'));rgba=(*rgba[:3],rgba[3]*float(s.get('stop-opacity',1))*opacity)
            stops.append((offset,rgba))
        if not stops:raise DisplayError('Empty gradient')
        def sample(t):
            if t<=stops[0][0]:return stops[0][1]
            for (a,ca),(b,cb) in zip(stops,stops[1:]):
                if t<=b:
                    f=(t-a)/(b-a) if b>a else 1
                    return tuple(x+(y-x)*f for x,y in zip(ca,cb))
            return stops[-1][1]
        # Resolve the target clip once; all bands stay within its visible shape.
        target=intersect(p,clip) if clip is not None else p
        if not len(target):return
        stats['gradient_objects']+=1
        inverse=Matrix(*gm);inverse.inverse()
        local=target.transform(*matrix(inverse));x0,y0,x1,y1=local.bounds
        extent=max(abs(v) for v in (x0,y0,x1,y1))+10
        def band(shape, rgba):append(intersect(shape.transform(*gm),target),rgba,None)
        if g.tag.endswith('linearGradient'):
            ax,ay,bx,by=[float(g.get(k,d)) for k,d in [('x1',0),('y1',0),('x2',1),('y2',0)]]
            dx,dy=bx-ax,by-ay;length=math.hypot(dx,dy)
            if length==0:append(target,stops[-1][1],None);return
            axis=(dx,dy,-dy/length,dx/length,ax,ay)
            # Make strips in gradient coordinates, extend before and after endpoint stops.
            bound=extent*4/length+4
            ranges=[(-bound,0)]+[(i/gradient_steps,(i+1)/gradient_steps) for i in range(gradient_steps)]+[(1,bound)]
            for lo,hi in ranges:
                shape=rectangle(lo,-extent*4,hi-lo,extent*8).transform(*axis)
                band(shape,sample((lo+hi)/2))
        elif g.tag.endswith('radialGradient'):
            cx,cy,fx,fy,r=[float(g.get(k,g.get(fallback,d))) for k,fallback,d in [('cx','cx',.5),('cy','cy',.5),('fx','cx',.5),('fy','cy',.5),('r','r',.5)]]
            if math.hypot(cx-fx,cy-fy)>r or g.get('fr') not in (None,'0'):
                raise DisplayError('Radial focal point requires explicit vector preparation')
            def disk(t):
                x=fx+(cx-fx)*t;y=fy+(cy-fy)*t;rr=r*t;k=.5522847498307936
                q=pathops.Path();q.moveTo(x+rr,y)
                q.cubicTo(x+rr,y+k*rr,x+k*rr,y+rr,x,y+rr)
                q.cubicTo(x-k*rr,y+rr,x-rr,y+k*rr,x-rr,y)
                q.cubicTo(x-rr,y-k*rr,x-k*rr,y-rr,x,y-rr)
                q.cubicTo(x+k*rr,y-rr,x+rr,y-k*rr,x+rr,y);q.close();return q
            outer=disk(1).transform(*gm)
            append(pathops.op(target,outer,pathops.PathOp.DIFFERENCE),sample(1),None)
            opaque=all(s[1][3]>=.999999 for s in stops)
            for i in reversed(range(gradient_steps)) if opaque else range(gradient_steps):
                hi=(i+1)/gradient_steps;lo=i/gradient_steps
                ring=disk(hi)
                # Opaque nested discs prevent hairline gaps from Office antialiasing.
                if i and not opaque:ring=pathops.op(ring,disk(lo),pathops.PathOp.DIFFERENCE)
                band(ring,sample((hi+lo)/2))
        else:raise DisplayError('Unsupported vector paint: '+g.tag)

    def reference(value):
        m=re.fullmatch(r'url\(#([^)]*)\)',value or '')
        if not m or m[1] not in ids:raise DisplayError('Unresolved local SVG resource')
        return m[1]

    def walk(n, transform=IDENTITY, style=None, clip=viewport, opacity=1., stack=()):
        tag=n.tag.split('}')[-1]
        if tag=='defs':return
        st=dict(style or {});st.update(n.attrib)
        tm=multiply(transform,matrix(n.get('transform')))
        if n.get('clip-path'):
            resource=ids[reference(n.get('clip-path'))];cp=pathops.Path()
            for child in resource:
                path=geometry(child).transform(*multiply(tm,matrix(child.get('transform'))))
                cp=pathops.op(cp,path,pathops.PathOp.UNION)
            clip=intersect(clip,cp) if clip is not None else cp
        opacity*=float(n.get('opacity',1))
        if n.get('mask'):
            mask=ids[reference(n.get('mask'))]
            rects=[v for v in mask.iter() if v.tag.endswith('}rect')]
            filters=[v for v in mask.iter() if v.get('filter')]
            if len(rects)!=1 or len(filters)!=1 or filters[0].get('filter')!='url(#filter-remove-color)':
                raise DisplayError('Complex mask needs explicit vector preparation')
            opacity*=float(rects[0].get('fill-opacity',1))
            stats['group_opacity_approximations']+=1
        if tag=='use':
            href=n.get('{http://www.w3.org/1999/xlink}href',n.get('href',''))
            if not href.startswith('#') or href[1:] not in ids or href in stack:raise DisplayError('Invalid SVG use')
            tm=multiply(tm,(1,0,0,1,float(n.get('x',0)),float(n.get('y',0))))
            if href.startswith('#glyph'):
                if clip is not None:
                    for glyph in ids[href[1:]].iter('{'+NS+'}path'):
                        glyph_path=geometry(glyph).transform(*tm)
                        outside=pathops.op(glyph_path,clip,pathops.PathOp.DIFFERENCE)
                        if len(outside) and abs(outside.area)>.01:
                            raise DisplayError('Partially clipped text needs editable text preparation; no outlining fallback')
                text=text_matcher.match(tm[4],tm[5])
                if text is not None:
                    text['color'][3]*=opacity
                    rows.append(text)
                return
            walk(ids[href[1:]],tm,st,clip,opacity,stack+(href,));return
        if tag in ('svg','g'):
            for child in n:walk(child,tm,st,clip,opacity,stack)
            return
        if tag in ('title','desc','metadata'):return
        if tag not in ('path','rect'):raise DisplayError('Unresolved SVG element: '+tag)
        original=geometry(n)
        for paint in ('fill','stroke'):
            value=st.get(paint,'black' if paint=='fill' else 'none')
            if value=='none':continue
            p=pathops.Path(original)
            if paint=='stroke':
                dash=[float(v) for v in re.split('[ ,]+',st.get('stroke-dasharray','').strip()) if v and v!='none']
                if len(dash)%2:dash*=2
                p.stroke(float(st.get('stroke-width',1)),
                         {'butt':pathops.LineCap.BUTT_CAP,'round':pathops.LineCap.ROUND_CAP,'square':pathops.LineCap.SQUARE_CAP}[st.get('stroke-linecap','butt')],
                         {'miter':pathops.LineJoin.MITER_JOIN,'round':pathops.LineJoin.ROUND_JOIN,'bevel':pathops.LineJoin.BEVEL_JOIN}[st.get('stroke-linejoin','miter')],
                         float(st.get('stroke-miterlimit',4)),dash or None,float(st.get('stroke-dashoffset',0)))
                p.convertConicsToQuads(.005)
            p=p.transform(*tm)
            alpha=opacity*float(st.get(paint+'-opacity',1))
            if value.startswith('url('):gradient(p,reference(value),tm,alpha,clip)
            else:
                rgba=color(value);append(p,(*rgba[:3],rgba[3]*alpha),clip)
    walk(root)
    for run in text_matcher.emitted:
        if any(i not in text_matcher.used for i in text_matcher.runs[run]):
            raise DisplayError('Partially clipped text needs explicit native text positioning; source retained')
    stats['native_text_runs']=sum(isinstance(row,dict) for row in rows)
    return rows,dims,stats,normalized


def child(parent, tag, **attributes):
    n=OxmlElement(tag)
    for k,v in attributes.items():n.set(k,str(v))
    parent.append(n);return n


def write_pptx(rows, dimensions, destination):
    prs=Presentation();w,h=dimensions['width_px'],dimensions['height_px']
    scale=min(9525.,51206400/max(w,h))
    prs.slide_width=round(w*scale);prs.slide_height=round(h*scale)
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    for index,row in enumerate(rows,2):
        if isinstance(row,dict):
            from native_text import add_textbox
            add_textbox(slide,row,scale,index)
            continue
        p,rgba=row
        if not len(p):continue
        x0,y0,x1,y1=p.bounds
        if x1-x0<1e-7 or y1-y0<1e-7:continue
        sp=OxmlElement('p:sp');nv=child(sp,'p:nvSpPr')
        child(nv,'p:cNvPr',id=index,name='Vector path '+str(index-1));child(nv,'p:cNvSpPr');child(nv,'p:nvPr')
        prop=child(sp,'p:spPr');tr=child(prop,'a:xfrm')
        child(tr,'a:off',x=round(x0*scale),y=round(y0*scale))
        ew,eh=max(1,round((x1-x0)*scale)),max(1,round((y1-y0)*scale))
        child(tr,'a:ext',cx=ew,cy=eh)
        geo=child(prop,'a:custGeom')
        for tag in ('avLst','gdLst','ahLst','cxnLst'):child(geo,'a:'+tag)
        child(geo,'a:rect',l=0,t=0,r='r',b='b')
        paths=child(geo,'a:pathLst');path=child(paths,'a:path',w=ew,h=eh,fill='norm',stroke='0',extrusionOk='0')
        for verb,pts in p:
            tag={pathops.PathVerb.MOVE:'moveTo',pathops.PathVerb.LINE:'lnTo',pathops.PathVerb.QUAD:'quadBezTo',pathops.PathVerb.CUBIC:'cubicBezTo',pathops.PathVerb.CLOSE:'close'}.get(verb)
            if tag is None:raise DisplayError('Unexpected conic path; cannot silently approximate')
            command=child(path,'a:'+tag)
            for x,y in pts:child(command,'a:pt',x=round((x-x0)*scale),y=round((y-y0)*scale))
        fill=child(prop,'a:solidFill');rgb=child(fill,'a:srgbClr',val=''.join(f'{max(0,min(255,round(v*255))):02X}' for v in rgba[:3]))
        if rgba[3]<1:child(rgb,'a:alpha',val=round(max(0,rgba[3])*100000))
        line=child(prop,'a:ln');child(line,'a:noFill')
        slide.shapes._spTree.insert_element_before(sp,'p:extLst')
    prs.save(str(destination))
    return len(slide.shapes)


def build(svg, output):
    output=Path(output);data=Path(svg).read_bytes()
    rows,dims,stats,normalized=native_paths(data)
    (output/'resolved.svg').write_bytes(normalized)
    pptx=output/'editable.pptx';count=write_pptx(rows,dims,pptx)
    return {'pptx':str(pptx),'native_freeforms':count-stats['native_text_runs'],'canvas':dims,'text_mode':'editable_text',**stats}
