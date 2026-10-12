"""Local native-path regressions; no Office, API, credentials or paid orders."""
import sys
import tempfile
import unittest
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from svg_native_geometry import native_paths, write_pptx, color
from display_check import inspect_pptx, DisplayError


def svg(body, viewbox='0 0 200 100'):
    return ('<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100" viewBox="'+viewbox+'">'+body+'</svg>').encode()


class NativePathsTests(unittest.TestCase):
    def test_nonzero_viewbox_nested_transforms_and_clip_are_resolved(self):
        data=svg('<defs><clipPath id="c"><rect x="0" y="0" width="8" height="5"/></clipPath></defs>'
                 '<g transform="translate(20 10)"><g transform="scale(2)" clip-path="url(#c)">'
                 '<rect x="-20" y="-20" width="80" height="80" fill="red"/></g></g>', '-10 -5 200 100')
        rows,_,_,_=native_paths(data)
        self.assertEqual(len(rows),1)
        for actual,expected in zip(rows[0][0].bounds,(30,15,46,25)):self.assertAlmostEqual(actual,expected,places=3)

    def test_evenodd_hole_and_opacity_survive(self):
        rows,_,_,_=native_paths(svg('<path d="M10 10H90V90H10Z M30 30H70V70H30Z" fill="red" fill-rule="evenodd" opacity=".25"/>'))
        self.assertFalse(rows[0][0].contains((50,50)))
        self.assertTrue(rows[0][0].contains((20,20)))
        self.assertAlmostEqual(rows[0][1][3],.25)
        self.assertEqual(color('white')[3],1)

    def test_gradient_is_many_editable_colors_without_media(self):
        data=svg('<defs><radialGradient id="r"><stop stop-color="red"/><stop offset="1" stop-color="blue"/></radialGradient></defs><circle cx="50" cy="50" r="40" fill="url(#r)"/>')
        rows,dims,stats,_=native_paths(data,gradient_steps=32)
        self.assertGreater(len(set(rgba for _,rgba in rows)),20)
        self.assertEqual(stats['gradient_objects'],1)
        with tempfile.TemporaryDirectory() as temp:
            ppt=Path(temp)/'out.pptx';write_pptx(rows,dims,ppt)
            audit=inspect_pptx(ppt)
            self.assertGreater(audit['native_freeforms'],20)
            self.assertEqual(audit['svg_graphics']+audit['raster_pictures'],0)
            with zipfile.ZipFile(ppt) as z:self.assertFalse(any(n.startswith('ppt/media/') for n in z.namelist()))

    def test_curves_round_caps_dashes_are_filled_paths(self):
        rows,_,_,_=native_paths(svg('<path d="M20 30C50 10 70 80 150 50" fill="none" stroke="green" stroke-width="6" stroke-linecap="round" stroke-dasharray="12 5"/>'))
        self.assertTrue(rows)
        self.assertTrue(all(p.bounds[2]<=155 for p,_ in rows))
        self.assertTrue(any(any(int(verb) in (2,4) for verb,_ in p) for p,_ in rows))

    def test_multiline_baselines_are_separate(self):
        rows,_,_,_=native_paths(svg('<text x="10" y="25" font-family="Times New Roman" font-size="12"><tspan x="10" dy="0">AAA</tspan><tspan x="10" dy="30">BBB</tspan></text>'))
        self.assertEqual([r['text'] for r in rows],['AAA','BBB'])
        self.assertAlmostEqual(rows[1]['baseline'][1]-rows[0]['baseline'][1],30,places=3)

    def test_subscript_keeps_later_text_baseline(self):
        normal,_,_,_=native_paths(svg('<text x="10" y="40" font-size="20">H<tspan font-size="12">2</tspan>O</text>'))
        shifted,_,stats,_=native_paths(svg('<text x="10" y="40" font-size="20">H<tspan font-size="12" baseline-shift="sub">2</tspan>O</text>'))
        self.assertEqual(stats['baseline_adjustments'],1)
        self.assertAlmostEqual(shifted[1]['baseline'][1]-normal[1]['baseline'][1],2.4,places=1)
        self.assertAlmostEqual(shifted[-1]['baseline'][1],normal[-1]['baseline'][1],places=1)

    def test_text_is_unicode_native_text_and_keeps_paint_order(self):
        data=svg('<rect width="40" height="40" fill="red"/><text x="10" y="25" font-family="Times New Roman" font-size="16">Cell 1</text><rect x="80" width="40" height="40" fill="blue"/>')
        rows,dims,stats,_=native_paths(data)
        self.assertEqual(rows[1]['text'],'Cell 1')
        self.assertFalse(stats['outlined_text'])
        with tempfile.TemporaryDirectory() as temp:
            ppt=Path(temp)/'text.pptx';write_pptx(rows,dims,ppt)
            from pptx import Presentation
            shapes=Presentation(ppt).slides[0].shapes
            self.assertEqual(len(shapes),3)
            text=shapes[1];self.assertEqual(text.text,'Cell 1')
            self.assertFalse(text.text_frame.word_wrap)
            self.assertEqual(text.text_frame.paragraphs[0].runs[0].font.name,'Times New Roman')
            self.assertEqual(inspect_pptx(ppt)['counts_by_type']['17'],1)

    def test_rotated_nested_text_retains_baseline_and_font_scale(self):
        data=svg('<g transform="translate(20 10) scale(2)"><text transform="rotate(20 20 20)" x="20" y="20" font-size="12">Ab</text></g>')
        rows,dims,_,_=native_paths(data)
        self.assertEqual(rows[0]['text'],'Ab')
        self.assertAlmostEqual(rows[0]['baseline'][0],60,places=1)
        self.assertAlmostEqual(rows[0]['baseline'][1],50,places=1)
        with tempfile.TemporaryDirectory() as temp:
            ppt=Path(temp)/'rotated.pptx';write_pptx(rows,dims,ppt)
            from pptx import Presentation
            shape=Presentation(ppt).slides[0].shapes[0]
            self.assertAlmostEqual(shape.rotation,20,places=2)
            self.assertAlmostEqual(shape.text_frame.paragraphs[0].runs[0].font.size.pt,18,places=2)

    def test_explicit_character_positions_are_not_collapsed(self):
        rows,_,_,_=native_paths(svg('<text x="10 60 120" y="40" font-size="16">ABC</text>'))
        self.assertEqual([r['text'] for r in rows],['A','B','C'])
        self.assertEqual([round(r['baseline'][0]) for r in rows],[10,60,120])

    def test_partly_clipped_text_is_not_silently_outlined(self):
        with self.assertRaises(DisplayError):
            native_paths(svg('<defs><clipPath id="c"><rect x="0" y="0" width="15" height="100"/></clipPath></defs><text x="10" y="40" font-size="25" clip-path="url(#c)">Hello</text>'))

    def test_unsupported_filter_is_not_silently_deleted(self):
        with self.assertRaises(DisplayError):native_paths(svg('<defs><filter id="f"><feGaussianBlur stdDeviation="3"/></filter></defs><rect width="30" height="30" filter="url(#f)"/>'))


if __name__=='__main__':unittest.main()
