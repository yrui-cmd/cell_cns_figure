"""Capture SVG text after layout, retaining Unicode and native editable text."""
import io
import math

import cairocffi as cairo
from cairosvg.parser import Tree
from cairosvg.surface import SVGSurface
from cairosvg.helpers import size

from display_check import DisplayError


class TextContext:
    def __init__(self, context, surface):
        self.raw=context;self.surface=surface

    def __getattr__(self, name):return getattr(self.raw,name)

    def show_text(self, text):
        ctx=self.raw;surface=self.surface;node=surface.text_node
        if node is None:return ctx.show_text(text)
        x,y=ctx.get_current_point();tm=ctx.get_matrix();fm=ctx.get_font_matrix()
        family=(node.get('font-family') or 'Times New Roman').split(',')[0].strip(' \"\'')
        try:rgba=ctx.get_source().get_rgba()
        except Exception as exc:raise DisplayError('Native text needs a solid font color; source retained') from exc
        record={'char':text,'baseline':list(ctx.user_to_device(x,y)),
                'matrix':list(tm),'font_matrix':list(fm),'font_family':family,
                'bold':node.get('font-weight','normal') in ('bold','bolder') or str(node.get('font-weight','')).isdigit() and int(node['font-weight'])>=550,
                'italic':node.get('font-style') in ('italic','oblique'),
                'underline':'underline' in node.get('text-decoration',''),
                'letter_spacing':size(surface,node.get('letter-spacing')),
                'color':list(rgba),'advance':ctx.text_extents(text)[4],
                'space_advance':ctx.text_extents(' ')[4],
                'ascent':ctx.font_extents()[0], 'descent':ctx.font_extents()[1],
                'run':surface.run_id,'full_text':node.text,
                'ink_box':list(ctx.text_extents(text)[:4])}
        surface.text_records.append(record)
        ctx.show_text(text)


class CaptureSurface(SVGSurface):
    def __init__(self,*args,**kwargs):
        self.text_records=[];self.text_node=None;self.run_id=0
        super().__init__(*args,**kwargs)

    def draw(self,node):
        if not isinstance(self.context,TextContext):self.context=TextContext(self.context,self)
        previous=self.text_node;previous_run=self.run_id
        if node.tag in ('text','tspan','textPath'):
            self.text_node=node;self.run_id=len(self.text_records)+1
        else:self.text_node=None
        try:super().draw(node)
        finally:self.text_node=previous;self.run_id=previous_run


def capture(data):
    output=io.BytesIO()
    surface=CaptureSurface(Tree(bytestring=data),output,96)
    surface.finish()
    return output.getvalue(),surface.text_records


class TextMatcher:
    """Replace normalized Cairo glyph uses in paint order, without text duplication."""
    def __init__(self, records):
        self.records=records;self.used=set();self.emitted=set();self.runs={}
        for index,record in enumerate(records):self.runs.setdefault(record['run'],[]).append(index)
        self.by_position={}
        for i,r in enumerate(records):
            x,y=r['baseline'];self.by_position.setdefault((round(x),round(y)),[]).append(i)

    def match(self,x,y):
        indices=[]
        for dx in (-1,0,1):
            for dy in (-1,0,1):indices.extend(self.by_position.get((round(x)+dx,round(y)+dy),[]))
        candidates=[(math.hypot(self.records[i]['baseline'][0]-x,self.records[i]['baseline'][1]-y),i) for i in indices if i not in self.used]
        if not candidates:raise DisplayError('Cannot associate SVG glyph with editable source text')
        distance,index=min(candidates)
        if distance>.1:raise DisplayError('Text coordinate association exceeds tolerance')
        self.used.add(index);record=self.records[index];run=record['run']
        if run in self.emitted:return None
        items=[self.records[i] for i in self.runs[run]]
        first=items[0];a,b,c,d,_,_=first['matrix'];sx=math.hypot(a,b);sy=math.hypot(c,d)
        if sx<1e-8 or sy<1e-8:raise DisplayError('Degenerate text transform')
        # Normal paragraphs become whole editable text runs. Explicit per-character
        # rotations/positions remain separate text boxes instead of losing geometry.
        same=all(all(abs(u-v)<1e-5 for u,v in zip(r['matrix'][:4],first['matrix'][:4])) and
                 abs((r['baseline'][0]-first['baseline'][0])*b-(r['baseline'][1]-first['baseline'][1])*a)<.05 for r in items)
        if same:
            cursor=0.;letter_index=0
            for i,ch in enumerate(first['full_text'].lstrip()):
                if i:cursor+=first.get('letter_spacing',0)*sx
                if ch.isspace():cursor+=first['space_advance']*sx;continue
                if letter_index>=len(items):same=False;break
                item=items[letter_index];letter_index+=1
                actual=((item['baseline'][0]-first['baseline'][0])*a+(item['baseline'][1]-first['baseline'][1])*b)/sx
                if abs(actual-cursor)>.12:same=False;break
                cursor+=item['advance']*sx
        if not same:
            first=record;items=[record];text=record['char']
        else:
            self.emitted.add(run);text=first['full_text'].lstrip()
            if ''.join(text.split())!=''.join(r['char'] for r in items):
                raise DisplayError('SVG text content cannot be reconstructed without changing characters')
        result=dict(first);result.update(kind='text',text=text,letters=items)
        return result


def add_textbox(slide, row, scale, index):
    from pptx.util import Pt
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    a,b,c,d,_,_=row['matrix'];sx=math.hypot(a,b);sy=math.hypot(c,d)
    if abs(a*c+b*d)>1e-5*sx*sy or a*d-b*c<=0:
        raise DisplayError('Skewed/reflected text needs explicit editable-text preparation')
    if abs(sx-sy)>max(sx,sy)*.001:
        raise DisplayError('Nonuniform text scaling needs explicit editable-text preparation; do not outline text')
    angle=math.atan2(b,a);co,si=math.cos(angle),math.sin(angle)
    font_size=abs(row['font_matrix'][3])*sy
    ascent=row['ascent']*sy
    first,last=row['letters'][0],row['letters'][-1]
    dx=last['baseline'][0]-first['baseline'][0];dy=last['baseline'][1]-first['baseline'][1]
    advance=dx*co+dy*si+last['advance']*sx
    width=max(advance+font_size*.15,font_size*.5);height=font_size*1.8
    x,y=row['baseline']
    # PowerPoint rotates about the center; locate that center from the SVG baseline.
    center_x=x+co*width/2-si*(height/2-ascent)
    center_y=y+si*width/2+co*(height/2-ascent)
    shape=slide.shapes.add_textbox(round((center_x-width/2)*scale),round((center_y-height/2)*scale),round(width*scale),round(height*scale))
    shape.name='Text '+str(index-1)+': '+row['text'][:60]
    shape._element.nvSpPr.cNvPr.set('id',str(index))
    shape.rotation=math.degrees(angle)%360
    tf=shape.text_frame;tf.clear();tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    p=tf.paragraphs[0];p.alignment=PP_ALIGN.LEFT
    p.space_before=p.space_after=Pt(0)
    run=p.add_run();run.text=row['text'];font=run.font
    family={'sans-serif':'Arial','serif':'Times New Roman','monospace':'Courier New'}.get(row['font_family'],row['font_family'])
    font.name=family;font.size=round(font_size*scale);font.bold=row['bold'];font.italic=row['italic'];font.underline=row.get('underline',False)
    font.color.rgb=RGBColor(*[max(0,min(255,round(v*255))) for v in row['color'][:3]])
    props=run._r.get_or_add_rPr();props.set('kern','400000')
    props.set('spc',str(round(row.get('letter_spacing',0)*sx*scale/12700*100)))
    for tag in ('a:ea','a:cs'):
        node=OxmlElement(tag);node.set('typeface',family);props.append(node)
    if row['color'][3]<1:
        alpha=OxmlElement('a:alpha');alpha.set('val',str(round(max(0,row['color'][3])*100000)))
        props.find('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')[0].append(alpha)
    # Store target baseline/advance for a reproducible local PowerPoint layout check.
    import json
    shape._element.nvSpPr.cNvPr.set('descr',json.dumps({'cell_text':True,'baseline_px':row['baseline'],
        'advance_px':advance,'font_size_px':font_size,'ascent_px':ascent,'rotation':math.degrees(angle),
        'family':family,'scale_emu_per_px':scale},ensure_ascii=False))
    return shape


def calibrate_baselines(pptx, output):
    """Measure this installed PowerPoint's first-line offset for each font face.

    Text stays text. A private calibration deck supplies actual pixel evidence;
    correcting its measured baseline offset avoids guessed textbox top margins.
    """
    import json
    from pathlib import Path
    import xml.etree.ElementTree as ET
    import cairosvg
    from PIL import Image, ImageChops
    from pptx import Presentation
    from svg_native_geometry import native_paths, write_pptx
    from display_check import render_powerpoint

    pptx=Path(pptx);output=Path(output)
    prs=Presentation(str(pptx));boxes=[];faces=[]
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text:continue
            try:meta=json.loads(shape._element.nvSpPr.cNvPr.get('descr','{}'))
            except ValueError:continue
            if not meta.get('cell_text'):continue
            font=shape.text_frame.paragraphs[0].runs[0].font
            key=(font.name,bool(font.bold),bool(font.italic))
            if key not in faces:faces.append(key)
            boxes.append((shape,meta,key))
    if not boxes:return []
    output.mkdir(parents=True,exist_ok=False);offsets={};records=[]
    sample_size=96
    for start in range(0,len(faces),6):
        batch=faces[start:start+6];folder=output/str(start//6+1);folder.mkdir()
        root=ET.Element('svg',{'xmlns':'http://www.w3.org/2000/svg','width':'1200','height':str(200*len(batch))})
        for i,(family,bold,italic) in enumerate(batch):
            n=ET.SubElement(root,'text',{'x':'30','y':str(130+i*200),'font-size':str(sample_size),
                'font-family':family,'font-weight':'bold' if bold else 'normal','font-style':'italic' if italic else 'normal'})
            n.text='AgHpx012'
        data=ET.tostring(root);source=folder/'source.svg';source.write_bytes(data)
        rows,dims,_,_=native_paths(data);deck=folder/'calibration.pptx';write_pptx(rows,dims,deck)
        expected=folder/'svg.png';cairosvg.svg2png(bytestring=data,write_to=str(expected),background_color='white')
        actual=render_powerpoint(source,folder,1200,200*len(batch),pptx=deck)
        with Image.open(expected) as f:reference=f.convert('RGB')
        with Image.open(actual['preview']) as f:candidate=f.convert('RGB')
        for i,key in enumerate(batch):
            bounds=[]
            for im in (reference,candidate):
                region=im.crop((0,200*i,1200,200*(i+1)))
                mask=ImageChops.difference(region,Image.new('RGB',region.size,'white')).convert('L').point(lambda v:255 if v>100 else 0)
                bbox=mask.getbbox()
                if bbox is None:raise DisplayError('PowerPoint font calibration rendered no text')
                bounds.append(bbox)
            offset=(bounds[1][1]-bounds[0][1])/sample_size
            if abs(offset)>.5:raise DisplayError('PowerPoint font baseline calibration is inconsistent')
            offsets[key]=offset
            records.append({'family':key[0],'bold':key[1],'italic':key[2],
                            'offset_per_font_px':offset,'reference_ink_bbox':bounds[0],'powerpoint_ink_bbox':bounds[1]})
    for shape,meta,key in boxes:
        correction=offsets[key]*meta['font_size_px'];previous=meta.get('baseline_correction_px',0)
        delta=(correction-previous)*meta['scale_emu_per_px'];angle=math.radians(meta['rotation'])
        shape.left+=round(math.sin(angle)*delta);shape.top-=round(math.cos(angle)*delta)
        meta['baseline_correction_px']=correction
        shape._element.nvSpPr.cNvPr.set('descr',json.dumps(meta,ensure_ascii=False))
    prs.save(str(pptx))
    (output/'baselines.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
    return records
