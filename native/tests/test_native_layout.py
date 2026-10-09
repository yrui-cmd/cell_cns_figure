import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import prepare_geometry_cache as ppt
import prepare_illustrator_cache as ai


def svg(body):
    return ET.fromstring('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1000 800">' + body + '</svg>')


def bounds(atom):
    points = [p['a'] for s in atom['subpaths'] for p in s['points']]
    return tuple(round(v, 4) for v in (min(p[0] for p in points), min(p[1] for p in points), max(p[0] for p in points), max(p[1] for p in points)))


class LayoutTests(unittest.TestCase):
    def both(self, body):
        left = ppt.collect_atoms(svg(body))
        right = ai.collect_atoms(svg(body))
        self.assertEqual([bounds(a) for a in left if a['kind'] == 'path'], [bounds(a) for a in right if a['kind'] == 'path'])
        return left

    def test_two_assets_do_not_collapse_to_origin(self):
        atoms = self.both('''<svg x="200" y="100" width="100" height="100" viewBox="0 0 10 10"><rect width="10" height="10"/></svg>
        <svg x="600" y="400" width="200" height="200" viewBox="0 0 10 10"><rect width="10" height="10"/></svg>''')
        self.assertEqual([bounds(a) for a in atoms], [(200, 100, 300, 200), (600, 400, 800, 600)])

    def test_nonzero_origin_default_meet_and_group_transform(self):
        atoms = self.both('''<g transform="translate(20 30) scale(2)"><svg x="100" y="50" width="200" height="100" viewBox="10 20 50 50"><rect x="10" y="20" width="50" height="50"/></svg></g>''')
        self.assertEqual(bounds(atoms[0]), (320, 130, 520, 330))

    def test_none_alignment_percentages_and_double_nesting(self):
        atoms = self.both('''<svg x="10%" y="25%" width="40%" height="25%" viewBox="0 0 100 100" preserveAspectRatio="none"><svg x="50%" width="50%" height="100%" viewBox="0 0 10 10" preserveAspectRatio="none"><rect width="10" height="10"/></svg></svg>''')
        self.assertEqual(bounds(atoms[0]), (300, 200, 500, 400))

    def test_alignment_and_nested_live_text(self):
        for builder in (ppt, ai):
            atoms = builder.collect_atoms(svg('''<svg x="50" y="60" width="200" height="100" viewBox="0 0 50 50" preserveAspectRatio="xMaxYMin meet"><text x="10" y="20" font-size="5">Test</text></svg>'''))
            self.assertEqual(atoms[0]['text']['position'], [170, 100])
            self.assertEqual(atoms[0]['text']['fontSize'], 10)

    def test_hidden_and_zero_viewport_do_not_draw(self):
        atoms = self.both('''<svg width="0"><rect width="100" height="100"/></svg><svg display="none"><rect width="100" height="100"/></svg><rect x="500" y="500" width="10" height="10"/>''')
        self.assertEqual(len(atoms), 1)

    def test_real_viewport_clip_is_calculated_without_repositioning(self):
        for builder in (ppt, ai):
            atoms=builder.collect_atoms(svg('''<svg x="200" y="100" width="100" height="100" viewBox="0 0 200 100" preserveAspectRatio="xMidYMid slice"><rect width="200" height="100"/></svg>'''))
            self.assertEqual(bounds(atoms[0]),(200,100,300,200))
            self.assertTrue(atoms[0]['viewport_clipped'])

    def test_zero_width_stroke_does_not_create_chemical_background_frame(self):
        atoms=self.both('''<svg width="234" height="229" stroke="black"><rect width="100%" height="100%" fill="white" fill-opacity="0" stroke-width="0"/><line x1="10" y1="20" x2="40" y2="20"/></svg>''')
        self.assertEqual(len(atoms),1)
        self.assertEqual(bounds(atoms[0]),(10,20,40,20))

    def test_percentage_geometry_uses_local_viewport(self):
        atoms=self.both('''<svg x="200" y="100" width="300" height="200"><rect x="10%" y="25%" width="50%" height="50%"/></svg>''')
        self.assertEqual(bounds(atoms[0]),(230,150,380,250))

    def test_move_only_path_is_not_an_error(self):
        atoms=self.both('<path d="M60 21"/><rect width="10" height="10"/>')
        self.assertEqual(len(atoms),1)

    def test_clipped_stroke_is_an_outline_without_false_edge_line(self):
        atoms=self.both('''<svg x="100" y="100" width="100" height="100"><path d="M-20 50L120 50" fill="none" stroke="red" stroke-width="10"/></svg>''')
        self.assertEqual(bounds(atoms[0]),(100,145,200,155))
        self.assertTrue(atoms[0]['paintParts'][0]['filled'])
        self.assertFalse(atoms[0]['paintParts'][0]['stroked'])

    def test_open_svg_contour_still_has_native_fill(self):
        from pptx import Presentation
        import run_cell_ppt_ooxml as writer
        atoms=self.both('<path d="M10 10L30 10L20 30" fill="#ff0000"/>')
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
        writer.add_freeform(slide,{'subpaths':atoms[0]['subpaths']},atoms[0]['paintParts'][0],'open-fill',2,(1,0,0,0,0))
        self.assertEqual(str(slide.shapes[0].fill.fore_color.rgb),'FF0000')

    def test_visible_overflow_keeps_geometry(self):
        atoms = self.both('''<svg x="200" width="100" height="100" viewBox="0 0 200 100" overflow="visible" preserveAspectRatio="xMidYMid slice"><rect width="200" height="100"/></svg>''')
        self.assertEqual(bounds(atoms[0]), (150, 0, 350, 100))

    def test_isolated_bundle_real_ppt_and_ai_cache(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            scripts = root/'standalone/scripts'
            shutil.copytree(SCRIPTS, scripts, ignore=shutil.ignore_patterns('__pycache__'))
            source = root/'input.svg'
            tree = svg('''<svg x="200" y="100" width="100" height="100" viewBox="0 0 10 10"><rect id="asset_a" width="10" height="10" fill="#ff0000"/></svg><svg x="600" y="400" width="200" height="200" viewBox="0 0 10 10"><rect id="asset_b" width="10" height="10" fill="#0000ff"/></svg><text x="100" y="700">Editable</text>''')
            source.write_bytes(ET.tostring(tree))
            original = source.read_bytes()
            for entry, name, extra in [('run_from_svg.py', 'shibielujing1', []), ('run_illustrator.py', 'shibielujing2', ['--dry-run'])]:
                proc = subprocess.run([sys.executable, '-X', 'utf8', str(scripts/entry), '--input-svg', str(source), '--output-root', str(root/'output'), '--job-name', name, *extra], capture_output=True, text=True, encoding='utf-8')
                self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
            deck = Presentation(root/'output/shibielujing1/shibielujing1.pptx')
            shapes = list(deck.slides[0].shapes)
            self.assertEqual(len(shapes), 3)
            self.assertGreater(shapes[1].left, shapes[0].left)
            self.assertGreater(shapes[1].top, shapes[0].top)
            self.assertAlmostEqual(shapes[1].width / shapes[0].width, 2, places=4)
            cache = json.loads((root/'output/shibielujing2/.illustrator-cache/geometry-cache.json').read_text(encoding='utf-8'))
            self.assertEqual([bounds(a) for a in cache['atoms'] if a['kind']=='path'], [(200, 100, 300, 200), (600, 400, 800, 600)])
            self.assertEqual(source.read_bytes(), original)


if __name__ == '__main__':
    unittest.main()
